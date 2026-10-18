"""B02 — сравнение LLM-провайдеров на одних и тех же документах.

Пример:
    python -m b02.compare --docs data/b02_docs --providers yandex,gigachat --out out/b02
    python -m b02.compare --docs data/b02_docs --providers mock --out out/b02_test   # без ключей

--docs: папка, где каждый подкаталог (или отдельный файл) — комплект материалов одной разработки.
Результат: out/review.xlsx (автометрики + лист для ручной оценки), out/raw/*.json (сырые ответы).
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from b02.providers import PROVIDERS  # noqa: E402
from tools.card import QUOTE_OK, check_card, fields, llm_schema, load_schema, normalize_card, quote_score  # noqa: E402
from tools.doc_text import approx_tokens, set_text  # noqa: E402
from tools.prompt import SYSTEM, user_message  # noqa: E402


def run_one(provider, schema, text, mode):
    modes = ["full", "flat"] if mode == "auto" else [mode]
    last_err = None
    for m in modes:
        t0 = time.time()
        try:
            raw_text, usage, raw = provider.complete(SYSTEM, user_message(schema, text), llm_schema(schema, m))
            obj = json.loads(raw_text)
            return {"mode": m, "latency": time.time() - t0, "usage": usage, "raw": raw, "obj": obj, "error": None}
        except Exception as e:  # пробуем следующий режим
            last_err = f"{type(e).__name__}: {e}"
            if hasattr(e, "response") and getattr(e, "response", None) is not None:
                last_err += " | " + e.response.text[:300]
    return {"mode": modes[-1], "latency": None, "usage": {}, "raw": None, "obj": None, "error": last_err}


def tokens(usage, *keys):
    for k in keys:
        if usage.get(k) is not None:
            return int(usage[k])
    return 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--docs", required=True)
    ap.add_argument("--providers", default="yandex,gigachat")
    ap.add_argument("--mode", default="auto", choices=["auto", "full", "flat"])
    ap.add_argument("--schema", default=None)
    ap.add_argument("--out", default="out/b02")
    a = ap.parse_args()

    schema = load_schema(a.schema) if a.schema else load_schema()
    out = Path(a.out)
    (out / "raw").mkdir(parents=True, exist_ok=True)
    sets = sorted(p for p in Path(a.docs).iterdir()
                  if not p.name.startswith(".") and p.resolve() != out.resolve())
    provs = [PROVIDERS[n.strip()]() for n in a.providers.split(",")]

    results = {}  # (set, provider) -> dict
    texts = {}
    for s in sets:
        texts[s.name] = set_text(s)
    skipped = [s.name for s in sets if not texts[s.name][0]]
    sets = [s for s in sets if texts[s.name][0]]
    for name in skipped:
        print(f"[{name}] пропущен: {'; '.join(texts[name][1])}")
    for s in sets:
        text, warns = texts[s.name]
        print(f"[{s.name}] ~{approx_tokens(text)} ток. {'; '.join(warns)}")
        for p in provs:
            r = run_one(p, schema, text, a.mode)
            r["warnings"] = warns
            r["tokens_doc"] = approx_tokens(text)
            if r["obj"] is not None:
                card = normalize_card(schema, r["obj"])
                r["card"] = card
                r["errors"] = check_card(schema, card)
                r["scores"] = {k: quote_score(v["quote"], text) for k, v in card.items()}
            (out / "raw" / f"{s.name}__{p.name}.json").write_text(
                json.dumps({k: v for k, v in r.items() if k != "raw"} | {"raw": r["raw"]},
                           ensure_ascii=False, indent=2, default=str), encoding="utf-8")
            status = "ок" if r["obj"] is not None else "ОШИБКА " + str(r["error"])[:200]
            print(f"   {p.name} [{r['mode']}]: {status}")
            results[(s.name, p.name)] = r

    write_xlsx(out / "review.xlsx", schema, sets, provs, results)
    print(f"\nГотово: {out / 'review.xlsx'}")


HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F4E79")
WRAP = Alignment(wrap_text=True, vertical="top")
BAD = PatternFill("solid", fgColor="FCE4D6")


def write_xlsx(path, schema, sets, provs, results):
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    rv = wb.create_sheet("Оценка")

    # ---- лист ручной оценки: строка = комплект × поле, по 4 колонки на провайдера
    head = ["Комплект", "Поле", "Название", "Тип"]
    for p in provs:
        head += [f"{p.name}: значение", f"{p.name}: цитата", f"{p.name}: цитата найдена", f"{p.name}: оценка"]
    rv.append(head)
    for s in sets:
        for key, spec in fields(schema):
            row = [s.name, key, spec["title"], spec["x-kind"]]
            for p in provs:
                r = results[(s.name, p.name)]
                c = r.get("card", {}).get(key, {})
                v = c.get("value")
                v = "\n".join(f"{x['name']}: {x['value']}" if isinstance(x, dict) else str(x) for x in v) \
                    if isinstance(v, list) else v
                sc = r.get("scores", {}).get(key)
                row += [v, c.get("quote"), None if sc is None else round(sc), None]
            rv.append(row)
    for i, h in enumerate(head, 1):
        cell = rv.cell(1, i)
        cell.font, cell.fill, cell.alignment = HEAD, FILL, WRAP
        rv.column_dimensions[get_column_letter(i)].width = 14 if i <= 4 or "найдена" in h or "оценка" in h else 45
    for row in rv.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
        for j in range(len(provs)):
            sc = row[4 + 4 * j + 2]
            if isinstance(sc.value, (int, float)) and sc.value < QUOTE_OK:
                sc.fill = BAD
    rv.freeze_panes = "E2"
    n = rv.max_row

    # ---- сводка
    ws.append(["B02 — сравнение LLM-провайдеров на одинаковых материалах, промпте и схеме"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Оценка на листе «Оценка»: 1 — верно (в т.ч. верно оставлено пустым), 0 — неверно, "
               "x — выдуманный факт (значения нет в материалах). Пустая ячейка — не оценено."])
    ws.append([])
    metrics = ["Модель", "Режим схемы", "Комплектов обработано", "Ошибок вызова", "Ошибок формата (справочники/коды)",
               "Заполнено полей, доля", "Цитата найдена в исходнике (из заполненных)", "Среднее время ответа, с",
               "Токенов на вход, всего", "Токенов на выход, всего", "Стоимость прогона, ₽ (если заданы цены)",
               "Ручная оценка: доля верных", "Ручная оценка: выдуманных фактов"]
    ws.append(["Показатель"] + [p.name for p in provs])
    nf = len(fields(schema))
    for mi, m in enumerate(metrics):
        row = [m]
        for j, p in enumerate(provs):
            rs = [results[(s.name, p.name)] for s in sets]
            ok = [r for r in rs if r["obj"] is not None]
            filled = [sc for r in ok for k, v in r["card"].items() if v["value"] is not None
                      for sc in [r["scores"][k]]]
            tin = sum(tokens(r["usage"], "prompt_tokens", "input_tokens") for r in rs)
            tout = sum(tokens(r["usage"], "completion_tokens", "output_tokens") for r in rs)
            col = get_column_letter(5 + 4 * j + 3)
            pin = os.environ.get(f"PRICE_{p.name.upper()}_IN")    # ₽ за 1000 входных токенов
            pout = os.environ.get(f"PRICE_{p.name.upper()}_OUT")  # ₽ за 1000 выходных токенов
            val = [
                p.model,
                ", ".join(sorted({r["mode"] for r in ok})) or "—",
                f"{len(ok)} из {len(rs)}",
                len(rs) - len(ok),
                sum(len(r["errors"]) for r in ok),
                round(sum(1 for r in ok for v in r["card"].values() if v["value"] is not None) / (nf * len(ok)), 3) if ok else None,
                round(sum(1 for sc in filled if sc is not None and sc >= QUOTE_OK) / len(filled), 3) if filled else None,
                round(sum(r["latency"] for r in ok) / len(ok), 1) if ok else None,
                tin, tout,
                round(tin / 1000 * float(pin) + tout / 1000 * float(pout), 2) if pin and pout else "задать PRICE_*",
                f'=IFERROR(COUNTIF(Оценка!{col}2:{col}{n},1)/COUNTA(Оценка!{col}2:{col}{n}),"—")',
                f'=COUNTIF(Оценка!{col}2:{col}{n},"x")',
            ][mi]
            row.append(val)
        ws.append(row)
    ws.cell(4, 1).font = Font(bold=True)
    for i in range(1, len(provs) + 2):
        ws.cell(4, i).font, ws.cell(4, i).fill = HEAD, FILL
    ws.column_dimensions["A"].width = 48
    for j in range(len(provs)):
        ws.column_dimensions[get_column_letter(2 + j)].width = 22
    ws.append([])
    ws.append(["Ошибки и предупреждения"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for (sname, pname), r in results.items():
        msgs = ([r["error"]] if r["error"] else []) + r.get("errors", []) + r.get("warnings", [])
        for msg in msgs:
            ws.append([f"{sname} / {pname}", msg])
    wb.save(path)


if __name__ == "__main__":
    main()
