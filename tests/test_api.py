"""Сквозной сценарий пилота через API: приглашение → карточка ИИ → загрузка → извлечение →
правка → отправка → модерация → выгрузка; плюс ручная карточка и проверки ошибок.

    pytest -q tests
"""
import io
import os
import tempfile
import time
from pathlib import Path

import pytest

os.environ["DATA_DIR"] = tempfile.mkdtemp(prefix="tt_test_")
os.environ["LLM_PROVIDER"] = "demo"
os.environ["ADMIN_PASSWORD"] = "secret"

from fastapi.testclient import TestClient  # noqa: E402

from backend.main import app, make_invite  # noqa: E402

DEMO = Path(__file__).resolve().parent.parent / "demo" / "agrosensor"
ADMIN = ("admin", "secret")


@pytest.fixture(scope="module")
def client():
    if not (DEMO / "spec.docx").exists():
        from demo.make_demo import main
        main()
    with TestClient(app) as c:
        yield c


def wait_ready(client, card_id, timeout=30):
    t0 = time.time()
    while time.time() - t0 < timeout:
        card = client.get(f"/api/cards/{card_id}").json()
        if card["status"] != "processing":
            return card
        time.sleep(0.2)
    raise AssertionError("извлечение не завершилось")


def test_full_pilot(client):
    assert client.get("/api/admin/cards").status_code == 401
    url = make_invite("Тест", ["ai", "manual"])
    token = url.split("t=")[1]

    assert client.get("/api/cards").status_code == 401
    assert client.post("/api/session", json={"token": "нет"}).status_code == 404
    s = client.post("/api/session", json={"token": token}).json()
    assert s["plan"] == ["ai", "manual"] and s["cards"] == []

    card = client.post("/api/cards").json()
    assert card["mode"] == "ai" and card["status"] == "draft"
    cid = card["id"]

    # без файлов извлечение не запускается; неподдерживаемый формат отклоняется
    assert client.post(f"/api/cards/{cid}/extract").status_code == 409
    r = client.post(f"/api/cards/{cid}/files", files=[("files", ("a.exe", b"MZ", "application/octet-stream"))])
    assert r.status_code == 415

    files = [("files", (p.name, p.read_bytes(), "application/octet-stream")) for p in sorted(DEMO.iterdir())]
    assert len(client.post(f"/api/cards/{cid}/files", files=files).json()) == 2

    assert client.post(f"/api/cards/{cid}/extract").json()["status"] == "processing"
    card = wait_ready(client, cid)
    assert card["status"] == "ready", card["ai"]
    assert card["values"]["org_inn"] == "1234567890"
    assert card["ai"]["card"]["org_name"]["verified"] is True
    assert card["title"]

    # после извлечения файлы не меняются, повторный запуск запрещён
    assert client.post(f"/api/cards/{cid}/extract").status_code == 409

    # сохранение с ошибкой формата не блокируется, но возвращает problems
    v = dict(card["values"])
    v["org_inn"] = "123"
    assert "org_inn" in client.put(f"/api/cards/{cid}/values", json={"values": v}).json()["problems"]

    # отправка с незаполненными обязательными полями → 422 со списком
    r = client.post(f"/api/cards/{cid}/submit")
    assert r.status_code == 422
    probs = r.json()["problems"]
    assert "org_inn" in probs and "tags" in probs

    v.update(org_inn="1234567890", tags=["датчики", "почва", "lorawan"],
             problem="Неравномерный полив и позднее обнаружение пересыхания почвы.",
             advantages="Работа 3 года без замены батареи.",
             technologies="LoRaWAN, машинное обучение для прогноза влажности.")
    assert client.put(f"/api/cards/{cid}/values", json={"values": v}).json()["problems"] == {}

    events = [{"card_id": cid, "type": "card_opened", "ts": "2026-10-18T10:00:00Z"},
              {"card_id": cid, "type": "submitted", "ts": "2026-10-18T10:12:30Z"},
              {"card_id": 99999, "type": "card_opened", "ts": "2026-10-18T10:00:00Z"}]
    assert client.post("/api/events", json=events).status_code == 204
    assert client.post(f"/api/cards/{cid}/submit").json()["status"] == "submitted"
    assert client.put(f"/api/cards/{cid}/values", json={"values": v}).status_code == 409

    # вторая карточка — ручная; третью план не допускает
    manual = client.post("/api/cards").json()
    assert manual["mode"] == "manual"
    assert client.post(f"/api/cards/{manual['id']}/extract").status_code == 409
    assert client.post("/api/cards").status_code == 409

    # модерация
    assert [c["id"] for c in client.get("/api/admin/cards?status=submitted", auth=ADMIN).json()] == [cid]
    a = client.get(f"/api/admin/cards/{cid}", auth=ADMIN).json()
    assert a["minutes"] == 12.5 and a["org_label"] == "Тест"
    assert client.get(f"/api/admin/cards/{cid}/files/{a['files'][0]['id']}", auth=ADMIN).status_code == 200
    r = client.post(f"/api/admin/cards/{cid}/moderation", auth=ADMIN, json={"verdict": "returned", "checklist": {}})
    assert r.status_code == 422  # возврат без комментария
    r = client.post(f"/api/admin/cards/{cid}/moderation", auth=ADMIN,
                    json={"verdict": "returned", "checklist": {"Наименование": True, "ТТХ": False},
                          "comment": "Добавьте единицы измерения"})
    assert r.json()["status"] == "returned"

    # организация правит возвращённую и отправляет повторно → принимаем
    assert client.post(f"/api/cards/{cid}/submit").json()["status"] == "submitted"
    client.post("/api/events", json=[{"card_id": cid, "type": "submitted", "ts": "2026-10-18T11:40:00Z"}])
    client.post(f"/api/admin/cards/{cid}/moderation", auth=ADMIN, json={"verdict": "accepted", "checklist": {}})
    a = client.get(f"/api/admin/cards/{cid}", auth=ADMIN).json()
    assert a["status"] == "accepted" and [m["verdict"] for m in a["moderation"]] == ["returned", "accepted"]
    assert a["minutes"] == 12.5  # доработка после возврата не входит во время заполнения

    # выгрузка: первая модерация — возврат → «принята с первого раза: нет»
    from openpyxl import load_workbook
    r = client.get("/api/admin/export.xlsx", auth=ADMIN)
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Карточки", "Поля ИИ", "События"]
    head = [c.value for c in wb["Карточки"][1]]
    row = {h: c.value for h, c in zip(head, wb["Карточки"][2])}
    assert row["минут на заполнение"] == 12.5 and row["принята с первого раза"] == "нет"
    ai_rows = list(wb["Поля ИИ"].iter_rows(min_row=2, values_only=True))
    assert any(r[1] == "org_inn" and r[4] == "нет" for r in ai_rows)


def test_extraction_error_returns_to_draft(client, monkeypatch):
    url = make_invite("Ошибка", ["ai"])
    client.post("/api/session", json={"token": url.split("t=")[1]})
    cid = client.post("/api/cards").json()["id"]
    client.post(f"/api/cards/{cid}/files", files=[("files", ("empty.txt", b"   \n ", "text/plain"))])
    client.post(f"/api/cards/{cid}/extract")
    card = wait_ready(client, cid)
    assert card["status"] == "draft" and card["ai"]["errors"]


def test_hallucination_is_flagged(client, monkeypatch):
    """mock возвращает организацию, которой нет в материалах, — цитата не подтверждается."""
    monkeypatch.setenv("LLM_PROVIDER", "mock")
    url = make_invite("Mock", ["ai"])
    client.post("/api/session", json={"token": url.split("t=")[1]})
    cid = client.post("/api/cards").json()["id"]
    client.post(f"/api/cards/{cid}/files", files=[("files", ("spec.docx", (DEMO / "spec.docx").read_bytes(), "x"))])
    client.post(f"/api/cards/{cid}/extract")
    card = wait_ready(client, cid)
    assert card["ai"]["provider"] == "mock"
    assert card["ai"]["card"]["org_name"]["verified"] is False
    assert card["ai"]["card"]["product_name"]["verified"] is True


def test_static_spa_fallback(client):
    assert client.get("/api/health").json()["ok"]
    assert client.get("/api/nope").status_code == 404
