"""C04/C05 — эталонная разметка в Excel и её перевод в JSON.

Пустой бланк на комплект (или с LLM-предразметкой из ответа B02/D02):
    python -m c04.gold template TT-001 --out data/gold_xlsx/TT-001.xlsx
    python -m c04.gold template TT-001 --prefill out/b02/raw/TT-001__yandex.json --out ...

Проверка и перевод заполненного бланка в эталон (с проверкой цитат по исходникам):
    python -m c04.gold import data/gold_xlsx/TT-001.xlsx --source data/raw/TT-001 --out data/gold/TT-001.json
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from tools.card import (KIND_RU, QUOTE_OK, check_card, fields, format_hint, load_schema,  # noqa: E402
                        normalize_card, quote_score)
from tools.doc_text import set_text  # noqa: E402

COLS = ["Раздел", "Поле", "Название", "Обяз.", "Как заполнять", "Формат / допустимые значения",
        "Значение", "Цитата из материалов", "Где (файл, стр./слайд)", "Комментарий"]
WIDTH = [24, 20, 26, 9, 15, 40, 45, 55, 20, 30]
FIRST = 8  # первая строка полей
HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F4E79")
INPUT = PatternFill("solid", fgColor="FFF9E5")
WRAP = Alignment(wrap_text=True, vertical="top")


def _cell_value(v):
    if v is None:
        return None
    if isinstance(v, list):
        return "\n".join(f"{x['name']}: {x['value']}" if isinstance(x, dict) else str(x) for x in v)
    return v


def template(set_id: str, out: Path, prefill: Path | None, schema: dict):
    pre = {}
    if prefill:
        raw = json.loads(prefill.read_text(encoding="utf-8"))
        pre = normalize_card(schema, raw.get("card") or raw.get("obj") or raw.get("fields") or raw)
    wb = Workbook()
    ws = wb.active
    ws.title = "Разметка"
    ws["A1"] = f"Эталонная карточка — шаблон {schema.get('x-version', '')}"
    ws["A1"].font = Font(bold=True, size=13)
    meta = [("Комплект (set_id)", set_id), ("Разметчик", None), ("Дата", dt.date.today().isoformat()),
            ("Время разметки, мин", None), ("Предразметка LLM", "да" if prefill else "нет")]
    for i, (k, v) in enumerate(meta, 2):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.cell(i, 2, v).fill = INPUT
    ws.cell(FIRST - 1, 1)  # заголовок таблицы
    for j, h in enumerate(COLS, 1):
        c = ws.cell(FIRST - 1, j, h)
        c.font, c.fill, c.alignment = HEAD, FILL, WRAP
        ws.column_dimensions[c.column_letter].width = WIDTH[j - 1]

    # справочники на скрытом листе — для выпадающих списков
    ref = wb.create_sheet("Справочники")
    ref.sheet_state = "hidden"
    col = 1
    for i, (key, spec) in enumerate(fields(schema)):
        r = FIRST + i
        hint = f"выбрать из выпадающего списка ({len(spec['enum'])} значений)" if "enum" in spec \
            else format_hint(spec)
        if spec["x-match"] in ("list", "table"):
            hint += ". Каждый элемент — с новой строки (Alt+Enter) или через « | »"
        p = pre.get(key, {})
        row = [spec["x-section"], key, spec["title"], "да" if spec["x-mvp"] == "required" else "",
               KIND_RU[spec["x-kind"]], f"{spec['description']}. {hint}",
               _cell_value(p.get("value")), p.get("quote"), None, None]
        for j, v in enumerate(row, 1):
            c = ws.cell(r, j, v)
            c.alignment = WRAP
            if j >= 7:
                c.fill = INPUT
        if "enum" in spec:
            for k, val in enumerate(spec["enum"], 1):
                ref.cell(k, col, val)
            letter = ref.cell(1, col).column_letter
            dv = DataValidation(type="list", formula1=f"=Справочники!${letter}$1:${letter}${len(spec['enum'])}",
                                allow_blank=True, showErrorMessage=True, errorTitle="Нет в справочнике",
                                error="Выберите значение из списка или оставьте пустым")
            ws.add_data_validation(dv)
            dv.add(ws.cell(r, 7))
            col += 1
    ws.freeze_panes = ws.cell(FIRST, 4)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Бланк: {out}")


def read_xlsx(path: Path, schema: dict) -> tuple[dict, dict]:
    ws = load_workbook(path, data_only=True)["Разметка"]
    meta = {ws.cell(i, 1).value: ws.cell(i, 2).value for i in range(2, FIRST - 1)}
    keys = dict(fields(schema))
    raw, extra = {}, {}
    for row in ws.iter_rows(min_row=FIRST, values_only=True):
        key = row[1]
        if key in keys:
            raw[key] = {"value": row[6], "quote": row[7]}
            extra[key] = {"where": row[8], "comment": row[9]}
    return meta, {"raw": raw, "extra": extra}


def import_gold(path: Path, out: Path, source: Path | None, schema: dict) -> int:
    meta, d = read_xlsx(path, schema)
    card = normalize_card(schema, d["raw"])
    problems = [f"ОШИБКА {e}" for e in check_card(schema, card)]
    text = set_text(source)[0] if source else None
    for key, v in card.items():
        if v["value"] is not None and not v["quote"]:
            problems.append(f"нет цитаты: {key}")
        if text and v["quote"]:
            sc = quote_score(v["quote"], text)
            if sc < QUOTE_OK:
                problems.append(f"цитата не найдена в материалах ({sc:.0f}): {key}")
        v.update({k: x for k, x in d["extra"][key].items() if x})
    gold = {"set_id": meta.get("Комплект (set_id)"), "schema_version": schema.get("x-version"),
            "annotator": meta.get("Разметчик"), "date": str(meta.get("Дата") or ""),
            "minutes": meta.get("Время разметки, мин"), "llm_prefill": meta.get("Предразметка LLM"),
            "fields": card}
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(gold, ensure_ascii=False, indent=2), encoding="utf-8")
    filled = sum(v["value"] is not None for v in card.values())
    print(f"{out}: заполнено {filled} из {len(card)} полей")
    for p in problems:
        print("  -", p)
    return sum(p.startswith("ОШИБКА") for p in problems)


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    t = sub.add_parser("template")
    t.add_argument("set_id")
    t.add_argument("--out", required=True)
    t.add_argument("--prefill")
    i = sub.add_parser("import")
    i.add_argument("xlsx")
    i.add_argument("--out", required=True)
    i.add_argument("--source", help="папка/файл с материалами комплекта — для проверки цитат")
    ap.add_argument("--schema")
    a = ap.parse_args()
    schema = load_schema(a.schema) if a.schema else load_schema()
    if a.cmd == "template":
        template(a.set_id, Path(a.out), Path(a.prefill) if a.prefill else None, schema)
    else:
        sys.exit(1 if import_gold(Path(a.xlsx), Path(a.out), Path(a.source) if a.source else None, schema) else 0)


if __name__ == "__main__":
    main()
