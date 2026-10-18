"""Технотека.ИИ — бэкенд прототипа (контракт: backend/openapi.yaml).

Запуск из корня репозитория:
    uvicorn backend.main:app --reload
Переменные окружения: см. .env.example.
"""
from __future__ import annotations

import hashlib
import io
import logging
import os
import re
import secrets
import shutil
import uuid
from contextlib import asynccontextmanager
from datetime import datetime, timezone
from pathlib import Path

from fastapi import BackgroundTasks, Depends, FastAPI, HTTPException, Request, Response, UploadFile
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.security import HTTPBasic, HTTPBasicCredentials
from fastapi.staticfiles import StaticFiles

from backend import db
from backend.db import jdump, jload, row, rows
from tools.card import check_card, fields, load_schema

log = logging.getLogger("tehnoteka")
logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")

SCHEMA = load_schema()
SCHEMA_VERSION = SCHEMA.get("x-version", "")
FIELD_KEYS = [k for k, _ in fields(SCHEMA)]
REQUIRED = SCHEMA.get("x-required-mvp", [])

ALLOWED_EXT = {".pdf", ".docx", ".pptx", ".html", ".htm", ".txt", ".md"}
MAX_FILES = int(os.environ.get("MAX_FILES", "5"))
MAX_FILE_MB = int(os.environ.get("MAX_FILE_MB", "20"))
ADMIN_USER = os.environ.get("ADMIN_USER", "admin")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "admin")
PUBLIC_URL = os.environ.get("PUBLIC_URL", "http://localhost:8000").rstrip("/")
COOKIE = "tt_session"
COOKIE_SECURE = PUBLIC_URL.startswith("https://")
EDITABLE = {"draft", "ready", "returned"}

@asynccontextmanager
async def lifespan(_app):
    _startup()
    yield


def _startup():
    db.init()
    if ADMIN_PASSWORD == "admin":
        log.warning("ADMIN_PASSWORD не задан — используется admin/admin. Только для локального запуска!")
    log.info("LLM_PROVIDER=%s, данные: %s", os.environ.get("LLM_PROVIDER", "mock"), db.DATA_DIR)


app = FastAPI(title="Технотека.ИИ", version="0.1.0", docs_url="/api/docs", openapi_url="/api/openapi.json",
              lifespan=lifespan)
basic = HTTPBasic(auto_error=False)


def now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")


def err(code: int, detail: str):
    raise HTTPException(code, detail)


# ---------------------------------------------------------------- доступ

def current_invite(request: Request) -> dict:
    token = request.cookies.get(COOKIE)
    if not token:
        err(401, "Нет сессии — откройте ссылку-приглашение")
    with db.connect() as c:
        inv = row(c, "SELECT * FROM invites WHERE token = ?", token)
    if not inv:
        err(401, "Приглашение не найдено")
    inv["plan"] = jload(inv["plan"], [])
    return inv


def admin(request: Request, creds: HTTPBasicCredentials | None = Depends(basic)) -> str:
    ok = creds is not None and \
        secrets.compare_digest(creds.username.encode(), ADMIN_USER.encode()) and \
        secrets.compare_digest(creds.password.encode(), ADMIN_PASSWORD.encode())
    if not ok:
        # фронт админки шлёт X-Requested-With — тогда без WWW-Authenticate, чтобы браузер не показывал своё окно
        hdr = {} if request.headers.get("x-requested-with") else {"WWW-Authenticate": "Basic"}
        raise HTTPException(401, "Неверный логин или пароль", headers=hdr)
    return creds.username


def own_card(card_id: int, inv: dict) -> dict:
    with db.connect() as c:
        card = row(c, "SELECT * FROM cards WHERE id = ? AND invite_id = ?", card_id, inv["id"])
    if not card:
        err(404, "Карточка не найдена")
    return card


# ---------------------------------------------------------------- представления

def title_of(values: dict) -> str | None:
    return values.get("product_short_name") or values.get("product_name") or None


def card_short(card: dict) -> dict:
    values = jload(card["values_json"], {})
    out = {"id": card["id"], "mode": card["mode"], "status": card["status"],
           "title": title_of(values), "updated_at": card["updated_at"].replace(" ", "T") + "Z"}
    if "org_label" in card:
        out["org_label"] = card["org_label"]
    return out


def card_full(c, card: dict) -> dict:
    files = rows(c, "SELECT id, name, size, mime FROM files WHERE card_id = ? ORDER BY id", card["id"])
    return {**card_short(card), "schema_version": card["schema_version"],
            "values": jload(card["values_json"], {}), "ai": jload(card["ai_json"]), "files": files}


def problems_of(values: dict) -> dict[str, str]:
    """Ошибки формата по полям — тем же check_card, что проверяет ответ модели."""
    probs = {}
    for p in check_card(SCHEMA, {k: {"value": v} for k, v in values.items() if k in FIELD_KEYS}):
        key, _, msg = p.partition(":")
        probs[key] = msg.strip()
    for k, v in values.items():
        spec = SCHEMA["properties"].get(k)
        if spec and spec["x-match"] == "table" and isinstance(v, list):
            if any(not str(r.get("name", "")).strip() or not str(r.get("value", "")).strip() for r in v):
                probs[k] = "в таблице есть незаполненные ячейки"
    return probs


def is_empty(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip()) or (isinstance(v, list) and not v)


def clean_values(values: dict) -> dict:
    out = {}
    for k, v in values.items():
        if k not in FIELD_KEYS or is_empty(v):
            continue
        if isinstance(v, str):
            v = v.strip()
        out[k] = v
    return out


# ---------------------------------------------------------------- пользователь

@app.get("/api/health")
def health():
    return {"ok": True, "provider": os.environ.get("LLM_PROVIDER", "mock"), "schema_version": SCHEMA_VERSION}


@app.get("/api/schema")
def get_schema():
    return SCHEMA


@app.post("/api/session")
def create_session(body: dict, response: Response):
    token = str(body.get("token", "")).strip()
    with db.connect() as c:
        inv = row(c, "SELECT * FROM invites WHERE token = ?", token)
        if not inv:
            err(404, "Приглашение не найдено или устарело")
        cards = rows(c, "SELECT * FROM cards WHERE invite_id = ? ORDER BY id", inv["id"])
    response.set_cookie(COOKIE, token, httponly=True, samesite="lax", secure=COOKIE_SECURE,
                        max_age=60 * 60 * 24 * 60)
    return {"org_label": inv["org_label"], "plan": jload(inv["plan"], []),
            "cards": [card_short(x) for x in cards]}


@app.get("/api/session")
def get_session(inv: dict = Depends(current_invite)):
    """Не в openapi v0.1: фронт проверяет, жива ли cookie, без повторного токена."""
    with db.connect() as c:
        cards = rows(c, "SELECT * FROM cards WHERE invite_id = ? ORDER BY id", inv["id"])
    return {"org_label": inv["org_label"], "plan": inv["plan"], "cards": [card_short(x) for x in cards]}


@app.get("/api/cards")
def list_my_cards(inv: dict = Depends(current_invite)):
    with db.connect() as c:
        return [card_short(x) for x in rows(c, "SELECT * FROM cards WHERE invite_id = ? ORDER BY id", inv["id"])]


@app.post("/api/cards", status_code=201)
def create_card(inv: dict = Depends(current_invite)):
    with db.connect() as c:
        n = row(c, "SELECT COUNT(*) AS n FROM cards WHERE invite_id = ?", inv["id"])["n"]
        if n >= len(inv["plan"]):
            err(409, f"По приглашению предусмотрено карточек: {len(inv['plan'])} — все уже созданы")
        mode = inv["plan"][n]
        cur = c.execute("INSERT INTO cards (invite_id, mode, schema_version) VALUES (?, ?, ?)",
                        (inv["id"], mode, SCHEMA_VERSION))
        card = row(c, "SELECT * FROM cards WHERE id = ?", cur.lastrowid)
        return card_full(c, card)


@app.get("/api/cards/{card_id}")
def get_card(card_id: int, inv: dict = Depends(current_invite)):
    card = own_card(card_id, inv)
    with db.connect() as c:
        out = card_full(c, card)
        last = row(c, "SELECT verdict, comment FROM moderation WHERE card_id = ? ORDER BY id DESC LIMIT 1", card_id)
    # организация видит только комментарий модератора к возврату, без чек-листа
    out["return_comment"] = last["comment"] if card["status"] == "returned" and last else None
    return out


def _safe_name(name: str) -> str:
    name = Path(name or "file").name
    name = re.sub(r"[^\w.\- ()«»№]", "_", name, flags=re.U).strip() or "file"
    return name[-120:]


@app.post("/api/cards/{card_id}/files")
async def upload_files(card_id: int, files: list[UploadFile], inv: dict = Depends(current_invite)):
    card = own_card(card_id, inv)
    if card["status"] != "draft":
        err(409, "Файлы можно менять только до запуска автозаполнения")
    with db.connect() as c:
        have = row(c, "SELECT COUNT(*) AS n FROM files WHERE card_id = ?", card_id)["n"]
    if have + len(files) > MAX_FILES:
        err(413, f"Не больше {MAX_FILES} файлов на карточку")
    saved = []
    for f in files:
        ext = Path(f.filename or "").suffix.lower()
        if ext not in ALLOWED_EXT:
            err(415, f"{f.filename}: формат не поддерживается (нужен PDF, DOCX или PPTX)")
        data = await f.read()
        if len(data) > MAX_FILE_MB * 1024 * 1024:
            err(413, f"{f.filename}: больше {MAX_FILE_MB} МБ")
        if not data:
            err(415, f"{f.filename}: пустой файл")
        saved.append((f, data))
    with db.connect() as c:
        for f, data in saved:
            name = _safe_name(f.filename)
            rel = Path(str(card_id)) / uuid.uuid4().hex[:12] / name
            path = db.FILES_DIR / rel
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(data)
            c.execute("INSERT INTO files (card_id, name, path, mime, size, sha256) VALUES (?, ?, ?, ?, ?, ?)",
                      (card_id, f.filename or name, str(rel), f.content_type or "application/octet-stream",
                       len(data), hashlib.sha256(data).hexdigest()))
        c.execute("UPDATE cards SET updated_at = ? WHERE id = ?", (now(), card_id))
        return rows(c, "SELECT id, name, size, mime FROM files WHERE card_id = ? ORDER BY id", card_id)


@app.delete("/api/cards/{card_id}/files/{file_id}", status_code=204)
def delete_file(card_id: int, file_id: int, inv: dict = Depends(current_invite)):
    card = own_card(card_id, inv)
    if card["status"] != "draft":
        err(409, "Файлы можно менять только до запуска автозаполнения")
    with db.connect() as c:
        f = row(c, "SELECT * FROM files WHERE id = ? AND card_id = ?", file_id, card_id)
        if not f:
            err(404, "Файл не найден")
        shutil.rmtree((db.FILES_DIR / f["path"]).parent, ignore_errors=True)
        c.execute("DELETE FROM files WHERE id = ?", (file_id,))
    return Response(status_code=204)


def run_extraction(card_id: int) -> None:
    """Фоновая задача: файлы → ml.extract_card → values (только непустые) + ai_json."""
    from ml.extract import ExtractionResult, extract_card
    with db.connect() as c:
        paths = [db.FILES_DIR / f["path"] for f in rows(c, "SELECT path FROM files WHERE card_id = ? ORDER BY id", card_id)]
    try:
        res = extract_card(paths, schema=SCHEMA)
    except Exception as e:  # ключи не заданы, сеть, неожиданный ответ — карточка не должна зависнуть
        log.exception("extraction failed for card %s", card_id)
        res = ExtractionResult(card={}, errors=[f"{type(e).__name__}: {str(e)[:300]}"])
    with db.connect() as c:
        if res.ok:
            values = {k: v["value"] for k, v in res.card.items() if not is_empty(v["value"])}
            c.execute("UPDATE cards SET status='ready', values_json=?, ai_json=?, updated_at=? WHERE id=?",
                      (jdump(values), res.to_json(), now(), card_id))
        else:
            c.execute("UPDATE cards SET status='draft', ai_json=?, updated_at=? WHERE id=?",
                      (res.to_json(), now(), card_id))
    log.info("card %s: extraction %s in %.1fs (%s)", card_id, "ok" if res.ok else "failed", res.latency_s, res.provider)


@app.post("/api/cards/{card_id}/extract", status_code=202)
def start_extraction(card_id: int, bg: BackgroundTasks, inv: dict = Depends(current_invite)):
    card = own_card(card_id, inv)
    if card["mode"] != "ai":
        err(409, "Эта карточка заполняется вручную")
    if card["status"] != "draft":
        err(409, "Автозаполнение уже запускалось")
    with db.connect() as c:
        if not row(c, "SELECT COUNT(*) AS n FROM files WHERE card_id = ?", card_id)["n"]:
            err(409, "Сначала загрузите материалы")
        c.execute("UPDATE cards SET status='processing', updated_at=? WHERE id=?", (now(), card_id))
        card = row(c, "SELECT * FROM cards WHERE id = ?", card_id)
    bg.add_task(run_extraction, card_id)
    return card_short(card)


@app.put("/api/cards/{card_id}/values")
def save_values(card_id: int, body: dict, inv: dict = Depends(current_invite)):
    card = own_card(card_id, inv)
    if card["status"] not in EDITABLE:
        err(409, "Карточка отправлена — редактирование закрыто")
    values = clean_values(body.get("values") or {})
    with db.connect() as c:
        c.execute("UPDATE cards SET values_json=?, updated_at=? WHERE id=?", (jdump(values), now(), card_id))
    return {"problems": problems_of(values)}


@app.post("/api/cards/{card_id}/submit")
def submit_card(card_id: int, inv: dict = Depends(current_invite)):
    card = own_card(card_id, inv)
    if card["status"] not in EDITABLE:
        err(409, "Карточка уже отправлена")
    values = jload(card["values_json"], {})
    probs = problems_of(values)
    for k in REQUIRED:
        if is_empty(values.get(k)):
            probs.setdefault(k, "обязательное поле")
    if probs:
        return JSONResponse(status_code=422, content={"problems": probs})
    with db.connect() as c:
        c.execute("UPDATE cards SET status='submitted', submitted_at=?, updated_at=? WHERE id=?",
                  (now(), now(), card_id))
        return card_short(row(c, "SELECT * FROM cards WHERE id = ?", card_id))


EVENT_TYPES = {"card_opened", "files_uploaded", "extraction_started", "extraction_done", "field_changed", "submitted"}


@app.post("/api/events", status_code=204)
def post_events(events: list[dict], inv: dict = Depends(current_invite)):
    with db.connect() as c:
        mine = {r["id"] for r in rows(c, "SELECT id FROM cards WHERE invite_id = ?", inv["id"])}
        for e in events[:500]:
            if e.get("card_id") in mine and e.get("type") in EVENT_TYPES and e.get("ts"):
                c.execute("INSERT INTO events (card_id, type, field, client_ts) VALUES (?, ?, ?, ?)",
                          (e["card_id"], e["type"], e.get("field"), str(e["ts"])[:40]))
    return Response(status_code=204)


# ---------------------------------------------------------------- модератор

ADMIN_CARD_SQL = "SELECT cards.*, invites.org_label FROM cards JOIN invites ON invites.id = cards.invite_id"


@app.get("/api/admin/cards")
def admin_list_cards(status: str | None = None, _: str = Depends(admin)):
    with db.connect() as c:
        if status:
            rs = rows(c, ADMIN_CARD_SQL + " WHERE status = ? ORDER BY cards.id DESC", status)
        else:
            rs = rows(c, ADMIN_CARD_SQL + " ORDER BY cards.id DESC")
    return [card_short(x) for x in rs]


@app.get("/api/admin/cards/{card_id}")
def admin_get_card(card_id: int, _: str = Depends(admin)):
    with db.connect() as c:
        card = row(c, ADMIN_CARD_SQL + " WHERE cards.id = ?", card_id)
        if not card:
            err(404, "Карточка не найдена")
        mods = rows(c, "SELECT * FROM moderation WHERE card_id = ? ORDER BY id", card_id)
        timing = row(c, "SELECT minutes FROM card_timing WHERE card_id = ?", card_id)
        out = card_full(c, card)
    out["moderation"] = [{"moderator": m["moderator"], "checklist": jload(m["checklist_json"], {}),
                          "verdict": m["verdict"], "comment": m["comment"],
                          "created_at": m["created_at"].replace(" ", "T") + "Z"} for m in mods]
    out["minutes"] = round(timing["minutes"], 1) if timing and timing["minutes"] is not None else None
    return out


@app.get("/api/admin/cards/{card_id}/files/{file_id}")
def admin_file(card_id: int, file_id: int, _: str = Depends(admin)):
    with db.connect() as c:
        f = row(c, "SELECT * FROM files WHERE id = ? AND card_id = ?", file_id, card_id)
    if not f:
        err(404, "Файл не найден")
    return FileResponse(db.FILES_DIR / f["path"], filename=f["name"], media_type=f["mime"])


@app.post("/api/admin/cards/{card_id}/moderation")
def moderate_card(card_id: int, body: dict, user: str = Depends(admin)):
    verdict = body.get("verdict")
    if verdict not in ("accepted", "returned"):
        err(422, "verdict: accepted или returned")
    if verdict == "returned" and not (body.get("comment") or "").strip():
        err(422, "При возврате напишите комментарий для организации")
    with db.connect() as c:
        card = row(c, "SELECT * FROM cards WHERE id = ?", card_id)
        if not card:
            err(404, "Карточка не найдена")
        if card["status"] != "submitted":
            err(409, "Модерировать можно только отправленную карточку")
        c.execute("INSERT INTO moderation (card_id, moderator, checklist_json, verdict, comment) VALUES (?, ?, ?, ?, ?)",
                  (card_id, user, jdump(body.get("checklist") or {}), verdict, body.get("comment")))
        c.execute("UPDATE cards SET status=?, updated_at=? WHERE id=?", (verdict, now(), card_id))
        return card_short(row(c, "SELECT * FROM cards WHERE id = ?", card_id))


def list_invites() -> list[dict]:
    """Для backend.cli: приглашения со ссылками (в API не выносим — создаёт и смотрит команда из консоли)."""
    with db.connect() as c:
        invs = rows(c, "SELECT i.*, (SELECT COUNT(*) FROM cards WHERE invite_id = i.id) AS cards "
                       "FROM invites i ORDER BY id DESC")
    return [{"id": i["id"], "org_label": i["org_label"], "plan": jload(i["plan"], []), "cards": i["cards"],
             "url": f"{PUBLIC_URL}/?t={i['token']}", "created_at": i["created_at"]} for i in invs]


def make_invite(org_label: str, plan: list[str]) -> str:
    if not org_label.strip() or not plan or any(p not in ("ai", "manual") for p in plan):
        raise ValueError("Нужны название организации и порядок из значений ai / manual")
    token = secrets.token_urlsafe(16)
    with db.connect() as c:
        c.execute("INSERT INTO invites (token, org_label, plan) VALUES (?, ?, ?)", (token, org_label.strip(), jdump(plan)))
    return f"{PUBLIC_URL}/?t={token}"


def _cell(v):
    if isinstance(v, list):
        return "\n".join(f"{x.get('name')}: {x.get('value')}" if isinstance(x, dict) else str(x) for x in v)
    return v


@app.get("/api/admin/export.xlsx")
def export_xlsx(_: str = Depends(admin)):
    """Три листа: карточки + время + модерация (КПЭ 1–2), поля ИИ → итог (сколько правок), события."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    with db.connect() as c:
        cards = rows(c, ADMIN_CARD_SQL + " ORDER BY cards.id")
        timing = {r["card_id"]: r["minutes"] for r in rows(c, "SELECT * FROM card_timing")}
        first = {r["card_id"]: r["accepted_first"] for r in rows(c, "SELECT * FROM first_pass")}
        events = rows(c, "SELECT * FROM events ORDER BY card_id, id")

    wb = Workbook()
    ws = wb.active
    ws.title = "Карточки"
    head = ["id", "организация", "режим", "статус", "минут на заполнение", "принята с первого раза",
            "заполнено обязательных", "провайдер", "модель", "сек. на извлечение"] + FIELD_KEYS
    ws.append(head)
    ws2 = wb.create_sheet("Поля ИИ")
    ws2.append(["card_id", "поле", "значение ИИ", "итоговое значение", "изменено пользователем",
                "цитата", "совпадение цитаты", "цитата подтверждена"])
    for card in cards:
        values = jload(card["values_json"], {})
        ai = jload(card["ai_json"]) or {}
        mins = timing.get(card["id"])
        acc = first.get(card["id"])
        filled = sum(not is_empty(values.get(k)) for k in REQUIRED)
        ws.append([card["id"], card["org_label"], card["mode"], card["status"],
                   round(mins, 1) if mins is not None else None,
                   None if acc is None else ("да" if acc else "нет"),
                   f"{filled}/{len(REQUIRED)}", ai.get("provider"), ai.get("model"), ai.get("latency_s")]
                  + [_cell(values.get(k)) for k in FIELD_KEYS])
        for k, f in (ai.get("card") or {}).items():
            if is_empty(f.get("value")):
                continue
            ws2.append([card["id"], k, _cell(f["value"]), _cell(values.get(k)),
                        "да" if f["value"] != values.get(k) else "нет",
                        f.get("quote"), f.get("quote_score"), "да" if f.get("verified") else "нет"])
    ws3 = wb.create_sheet("События")
    ws3.append(["card_id", "тип", "поле", "время браузера", "время сервера"])
    for e in events:
        ws3.append([e["card_id"], e["type"], e["field"], e["client_ts"], e["server_ts"]])
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 18
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"tehnoteka_{datetime.now():%Y%m%d_%H%M}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})


# ---------------------------------------------------------------- фронтенд (собранный) из backend/static

STATIC = Path(os.environ.get("STATIC_DIR", Path(__file__).resolve().parent / "static"))
if (STATIC / "index.html").exists():
    app.mount("/assets", StaticFiles(directory=STATIC / "assets"), name="assets")

    @app.get("/{path:path}", include_in_schema=False)
    def spa(path: str):
        if path.startswith("api/"):
            err(404, "Нет такого метода")
        f = STATIC / path
        return FileResponse(f if path and f.is_file() else STATIC / "index.html")
