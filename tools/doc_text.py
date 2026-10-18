"""Материалы разработки → один текст для модели.

Комплект = папка с файлами одной разработки (или один файл).
Поддержка: PDF (pymupdf), DOCX (python-docx), PPTX (python-pptx), HTML (beautifulsoup4), TXT/MD.
В текст вставляются метки «--- стр. N ---» / «--- слайд N ---», чтобы разметчик и модель
могли сослаться на место цитаты. Сканы без текстового слоя не распознаются (OCR вне MVP).
"""
from __future__ import annotations

from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".pptx", ".txt", ".md", ".html", ".htm"}

# служебные блоки страницы, которые не относятся к описанию продукта
HTML_DROP = ["script", "style", "noscript", "svg", "nav", "header", "footer", "form", "iframe", "button"]


def _pdf(path: Path) -> str:
    import pymupdf
    with pymupdf.open(path) as doc:
        return "\n".join(f"--- стр. {i} ---\n{page.get_text()}" for i, page in enumerate(doc, 1))


def _docx(path: Path) -> str:
    from docx import Document
    d = Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = []
            for c in row.cells:  # объединённые ячейки python-docx повторяет — убираем дубли подряд
                if not cells or cells[-1] != c.text.strip():
                    cells.append(c.text.strip())
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def _pptx(path: Path) -> str:
    from pptx import Presentation
    out = []
    for i, slide in enumerate(Presentation(str(path)).slides, 1):
        out.append(f"--- слайд {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    out.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append("Заметки: " + slide.notes_slide.notes_text_frame.text)
    return "\n".join(out)


def _html(path: Path) -> str:
    """Сохранённая веб-страница → текст. Таблицы — строками «ячейка | ячейка»."""
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(path.read_bytes(), "html.parser")
    for tag in soup(HTML_DROP):
        tag.decompose()
    for row in soup.find_all("tr"):
        cells = [c.get_text(" ", strip=True) for c in row.find_all(["td", "th"])]
        row.replace_with(soup.new_string("\n" + " | ".join(c for c in cells if c) + "\n"))
    title = soup.title.get_text(strip=True) if soup.title else ""
    body = soup.body or soup
    lines = [ln.strip() for ln in body.get_text("\n").splitlines()]
    out, prev = [], None
    for ln in lines:  # убираем пустые строки и повторы подряд (меню, дубли блоков)
        if ln and ln != prev:
            out.append(ln)
        prev = ln
    return (f"Заголовок страницы: {title}\n" if title else "") + "\n".join(out)


def file_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _pdf(path)
    if ext == ".docx":
        return _docx(path)
    if ext == ".pptx":
        return _pptx(path)
    if ext in (".html", ".htm"):
        return _html(path)
    return path.read_text(encoding="utf-8", errors="replace")


def set_files(set_path: Path) -> list[Path]:
    if set_path.is_file():
        return [set_path] if set_path.suffix.lower() in SUPPORTED else []
    return sorted(p for p in set_path.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED
                  and not p.name.startswith(("~$", ".")))


def set_text(set_path: Path) -> tuple[str, list[str]]:
    """Текст комплекта с разделителями по файлам и список предупреждений."""
    parts, warnings = [], []
    for f in set_files(set_path):
        try:
            t = file_text(f).strip()
        except Exception as e:  # битый файл не должен ронять весь прогон
            warnings.append(f"{f.name}: не удалось прочитать ({type(e).__name__}: {e})")
            continue
        body = "\n".join(ln for ln in t.splitlines() if not ln.startswith("--- "))
        if len(body.strip()) < 200:
            warnings.append(f"{f.name}: почти нет текста ({len(body.strip())} симв.) — возможно, скан")
        parts.append(f"=== Файл: {f.name} ===\n{t}")
    if not parts:
        warnings.append("в комплекте нет читаемых файлов")
    return "\n\n".join(parts), warnings


def approx_tokens(text: str) -> int:
    """Грубая оценка для русского текста (~3,5 символа на токен)."""
    return int(len(text) / 3.5)
